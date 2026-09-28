"""Fetch KArSL Arabic alphabet classes 32-70 without downloading the 25 GB archive."""

from __future__ import annotations

import subprocess
from pathlib import Path

import requests
from openpyxl import load_workbook
from remotezip import RemoteZip

ROOT = Path(__file__).resolve().parents[1]
DOWNLOAD = "https://www.kaggle.com/api/v1/datasets/download/yousefdotpy/karsl-502?datasetVersionNumber=3"


def archive_url() -> str:
    response = requests.get(DOWNLOAD, headers={"Range": "bytes=0-0"}, allow_redirects=True, stream=True, timeout=60)
    response.raise_for_status()
    url = response.url
    response.close()
    return url


def main() -> None:
    sheet = load_workbook(
        ROOT / "references" / "karsl_word_recognition" / "KARSL-502_Labels.xlsx",
        read_only=True, data_only=True,
    ).active
    labels = {int(row[0]): str(row[1]) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] and 32 <= int(row[0]) <= 70}
    video_root = ROOT / "data" / "KArSL_alphabet"
    frame_root = ROOT / "data" / "KArSL_alphabet_frames"
    with RemoteZip(archive_url()) as archive:
        chosen: dict[int, list] = {key: [] for key in labels}
        sequence: dict[int, str] = {}
        for member in archive.infolist():
            parts = member.filename.split("/")
            if len(parts) < 6 or parts[:3] != ["01", "01", "test"] or not parts[3].isdigit() or not member.filename.lower().endswith(".jpg"):
                continue
            sign_id = int(parts[3])
            if sign_id not in chosen:
                continue
            sequence.setdefault(sign_id, parts[4])
            if parts[4] == sequence[sign_id]:
                chosen[sign_id].append(member)
        for count, (sign_id, members) in enumerate(chosen.items(), 1):
            destination = video_root / f"{sign_id:04d}"
            destination.mkdir(parents=True, exist_ok=True)
            (destination / "label.txt").write_text(labels[sign_id], encoding="utf-8")
            output = destination / "sign.mp4"
            if output.exists() and output.stat().st_size > 1000:
                print(f"[{count:02d}/39] {labels[sign_id]} ready", flush=True)
                continue
            local = frame_root / f"{sign_id:04d}"
            local.mkdir(parents=True, exist_ok=True)
            for index, member in enumerate(members, 1):
                target = local / f"{index:04d}.jpg"
                if not target.exists() or target.stat().st_size < 1000:
                    with archive.open(member) as source, target.open("wb") as sink:
                        sink.write(source.read())
            subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-framerate", "25", "-i", str(local / "%04d.jpg"),
                            "-c:v", "libx264", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(output)], check=True)
            print(f"[{count:02d}/39] {labels[sign_id]}: {len(members)} frames", flush=True)


if __name__ == "__main__":
    main()
