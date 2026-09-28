"""Build one representative H.264 output clip for every KArSL-502 class."""

from __future__ import annotations

import shutil
import struct
import subprocess
import zlib
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import requests
from openpyxl import load_workbook
from remotezip import RemoteZip

ROOT = Path(__file__).resolve().parents[1]
DOWNLOAD = "https://www.kaggle.com/api/v1/datasets/download/yousefdotpy/karsl-502?datasetVersionNumber=3"


def archive_url() -> str:
    response = requests.get(DOWNLOAD, headers={"Range": "bytes=0-0"}, allow_redirects=True, stream=True, timeout=60)
    response.raise_for_status()
    url = response.url
    response.close()
    return url


def labels() -> dict[int, str]:
    sheet = load_workbook(
        ROOT / "references" / "karsl_word_recognition" / "KARSL-502_Labels.xlsx",
        read_only=True, data_only=True,
    ).active
    return {int(row[0]): str(row[1]).strip() for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]}


def reuse_existing(sign_id: int, destination: Path) -> bool:
    sources = [
        ROOT / "data" / "KArSL_alphabet" / f"{sign_id:04d}" / "sign.mp4",
        ROOT / "data" / "KArSL100_dictionary" / f"{sign_id:04d}" / "sign.mp4",
    ]
    for source in sources:
        if source.exists() and source.stat().st_size > 1000:
            shutil.copy2(source, destination)
            return True
    return False


def main() -> None:
    class_labels = labels()
    video_root = ROOT / "data" / "KArSL502_dictionary"
    frame_root = ROOT / "data" / "KArSL502_dictionary_frames"
    video_root.mkdir(parents=True, exist_ok=True)

    # Reuse the 139 clips already downloaded before touching the network.
    for sign_id, label in class_labels.items():
        folder = video_root / f"{sign_id:04d}"
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "label.txt").write_text(label, encoding="utf-8")
        output = folder / "sign.mp4"
        if not output.exists():
            reuse_existing(sign_id, output)

    missing_ids = {
        sign_id for sign_id in class_labels
        if not (video_root / f"{sign_id:04d}" / "sign.mp4").exists()
    }
    print(f"Reused {len(class_labels) - len(missing_ids)}; fetching {len(missing_ids)} classes", flush=True)
    if not missing_ids:
        return

    signed_url = archive_url()
    with RemoteZip(signed_url) as archive:
        chosen: dict[int, list] = {key: [] for key in missing_ids}
        sequence: dict[int, str] = {}
        for member in archive.infolist():
            parts = member.filename.split("/")
            if len(parts) < 6 or parts[:3] != ["01", "01", "test"] or not parts[3].isdigit() or not member.filename.lower().endswith(".jpg"):
                continue
            sign_id = int(parts[3])
            if sign_id not in chosen:
                continue
            sequence.setdefault(sign_id, parts[4])
            if parts[4] == sequence[sign_id]:
                chosen[sign_id].append(member)

        def build_one(sign_id: int, members: list) -> tuple[int, int]:
            if not members:
                return sign_id, 0
            local = frame_root / f"{sign_id:04d}"
            local.mkdir(parents=True, exist_ok=True)
            # All JPEGs in one sequence are adjacent ZIP members. Fetch their
            # complete byte span once instead of making one HTTP request/frame.
            start = members[0].header_offset
            last = members[-1]
            end = last.header_offset + 30 + len(last.filename.encode("utf-8")) + last.compress_size + 4096
            response = requests.get(signed_url, headers={"Range": f"bytes={start}-{end}"}, timeout=120)
            response.raise_for_status()
            blob = response.content
            for index, member in enumerate(members, 1):
                target = local / f"{index:04d}.jpg"
                if not target.exists() or target.stat().st_size < 1000:
                    offset = member.header_offset - start
                    header = struct.unpack_from("<IHHHHHIIIHH", blob, offset)
                    if header[0] != 0x04034B50:
                        raise RuntimeError(f"Invalid ZIP member header: {member.filename}")
                    name_length, extra_length = header[-2], header[-1]
                    data_start = offset + 30 + name_length + extra_length
                    compressed = blob[data_start:data_start + member.compress_size]
                    content = zlib.decompress(compressed, -15) if member.compress_type == 8 else compressed
                    target.write_bytes(content)
            output = video_root / f"{sign_id:04d}" / "sign.mp4"
            subprocess.run([
                "ffmpeg", "-y", "-loglevel", "error", "-framerate", "25", "-i", str(local / "%04d.jpg"),
                "-c:v", "libx264", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(output),
            ], check=True)
            return sign_id, len(members)

        # Range reads are latency-bound; parallel class extraction is substantially faster.
        with ThreadPoolExecutor(max_workers=8) as executor:
            futures = {executor.submit(build_one, sign_id, members): sign_id for sign_id, members in chosen.items()}
            for count, future in enumerate(as_completed(futures), 1):
                sign_id, frame_count = future.result()
                print(f"[{count:03d}/{len(chosen)}] {sign_id:04d} {class_labels[sign_id]} ({frame_count} frames)", flush=True)


if __name__ == "__main__":
    main()
