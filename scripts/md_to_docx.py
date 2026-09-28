"""Convert project markdown docs to Word (.docx) using python-docx."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]


def add_runs_with_bold(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def convert(md_path: Path, docx_path: Path) -> None:
    base_dir = md_path.parent
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith("|") and "|" in line.strip()[1:]:
            row = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if row and not all(set(c) <= {"-", ":"} for c in row):
                table_rows.append(row)
            i += 1
            if i < len(lines) and lines[i].strip().startswith("|"):
                continue
            if table_rows:
                table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                table.style = "Table Grid"
                for r_idx, row in enumerate(table_rows):
                    for c_idx, cell in enumerate(row):
                        table.rows[r_idx].cells[c_idx].text = cell
                table_rows = []
            continue
        else:
            if table_rows:
                table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                table.style = "Table Grid"
                for r_idx, row in enumerate(table_rows):
                    for c_idx, cell in enumerate(row):
                        table.rows[r_idx].cells[c_idx].text = cell
                table_rows = []

        img = re.match(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if img:
            alt, rel = img.group(1), img.group(2)
            img_path = (base_dir / rel).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            if img_path.exists():
                doc.add_picture(str(img_path), width=Inches(6.0))
            else:
                p.add_run(f"[Image: {alt} — {rel}]")
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.strip() == "---":
            doc.add_paragraph("")
        elif line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, line.strip()[2:])
        elif re.match(r"^\d+\.\s", line.strip()):
            p = doc.add_paragraph(style="List Number")
            add_runs_with_bold(p, re.sub(r"^\d+\.\s", "", line.strip()))
        elif line.strip():
            p = doc.add_paragraph()
            add_runs_with_bold(p, line.strip())
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"Wrote {docx_path}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "docs")
    args = parser.parse_args()
    for md in args.inputs:
        out = args.output_dir / f"{md.stem}.docx"
        convert(md.resolve(), out.resolve())


if __name__ == "__main__":
    main()
