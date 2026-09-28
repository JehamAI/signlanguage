"""Inference adapter for the pretrained SignBart KArSL-502 skeleton model."""

from __future__ import annotations

import json
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import cv2
import mediapipe as mp
import numpy as np
import torch
from openpyxl import load_workbook
from safetensors.torch import load_file

from .config import ROOT, Settings


@dataclass(frozen=True)
class KArSL502Prediction:
    gloss: str
    confidence: float
    accepted: bool
    frames_with_hands: int
    sampled_frames: int


BODY = list(range(11, 17))
LEFT = list(range(33, 54))
RIGHT = list(range(54, 75))
PARTS = (BODY, LEFT, RIGHT)


def _labels() -> list[str]:
    workbook = ROOT / "references" / "karsl_word_recognition" / "KARSL-502_Labels.xlsx"
    sheet = load_workbook(workbook, read_only=True, data_only=True).active
    rows = sorted(
        ((int(row[0]), str(row[1]).strip()) for row in sheet.iter_rows(min_row=2, values_only=True)),
        key=lambda item: item[0],
    )
    if len(rows) != 502 or [item[0] for item in rows] != list(range(1, 503)):
        raise ValueError("KArSL-502 label workbook must contain IDs 1 through 502")
    return [item[1] for item in rows]


def _normalize_part(points: np.ndarray) -> np.ndarray:
    result = points.copy()
    min_x, min_y = result.min(axis=0)
    max_x, max_y = result.max(axis=0)
    width, height = max_x - min_x, max_y - min_y
    if width > height:
        delta_x = 0.05 * width
        delta_y = delta_x + ((width - height) / 2)
    else:
        delta_y = 0.05 * height
        delta_x = delta_y + ((height - width) / 2)
    start = np.maximum([min_x - delta_x, min_y - delta_y], 0.0)
    end = np.minimum([max_x + delta_x, max_y + delta_y], 1.0)
    scale = end - start
    if scale[0] != 0:
        result[:, 0] = (result[:, 0] - start[0]) / scale[0]
    if scale[1] != 0:
        result[:, 1] = (result[:, 1] - start[1]) / scale[1]
    return result


def _points(landmarks, count: int) -> np.ndarray:
    if landmarks is None:
        return np.zeros((count, 2), dtype=np.float32)
    return np.asarray([(point.x, point.y) for point in landmarks.landmark], dtype=np.float32)


class SignBartKArSL502Recognizer:
    def __init__(self, threshold: float = 0.60):
        from references.SignBart.model import SignBart

        artifact = ROOT / "artifacts" / "karsl502_signbart"
        config_path = artifact / "config.json"
        weights_path = artifact / "model.safetensors"
        if not config_path.exists() or not weights_path.exists():
            raise FileNotFoundError("The SignBart KArSL-502 checkpoint is missing")
        config = json.loads(config_path.read_text(encoding="utf-8"))
        config["num_labels"] = 502
        self.labels = _labels()
        self.threshold = threshold
        self.device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
        self.model = SignBart(config)
        state = load_file(str(weights_path))
        missing, unexpected = self.model.load_state_dict(state, strict=False)
        harmless = {"encoder.embed_tokens.weight", "decoder.embed_tokens.weight"}
        if missing or set(unexpected) - harmless:
            raise RuntimeError(f"SignBart checkpoint mismatch: missing={missing}, unexpected={unexpected}")
        self.model.to(self.device).eval()

    def _extract(self, video: Path, maximum_frames: int = 64) -> tuple[np.ndarray, int, int]:
        capture = cv2.VideoCapture(str(video))
        frames: list[np.ndarray] = []
        hand_frames = 0
        with mp.solutions.holistic.Holistic(
            static_image_mode=False,
            model_complexity=1,
            min_detection_confidence=0.35,
            min_tracking_confidence=0.35,
        ) as holistic:
            while True:
                ok, frame = capture.read()
                if not ok:
                    break
                result = holistic.process(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))
                pose = _points(result.pose_landmarks, 33)
                left = _points(result.left_hand_landmarks, 21)
                right = _points(result.right_hand_landmarks, 21)
                has_hands = result.left_hand_landmarks is not None or result.right_hand_landmarks is not None
                if not has_hands:
                    continue
                hand_frames += 1
                # The released checkpoint used the opposite handedness ordering from
                # MediaPipe's selfie convention. Swapping restores the training layout.
                left, right = right, left
                points = np.concatenate((pose, left, right), axis=0)
                for part in PARTS:
                    points[part] = _normalize_part(points[part])
                frames.append(points)
        capture.release()
        observed = len(frames)
        if not frames:
            return np.empty((0, 75, 2), dtype=np.float32), hand_frames, observed
        values = np.asarray(frames, dtype=np.float32)
        if len(values) > maximum_frames:
            indexes = np.linspace(0, len(values) - 1, maximum_frames).astype(int)
            values = values[indexes]
        return values, hand_frames, observed

    def predict_video(self, path: Path) -> KArSL502Prediction:
        values, hand_frames, observed = self._extract(path)
        if not len(values):
            return KArSL502Prediction("", 0.0, False, hand_frames, observed)
        keypoints = torch.from_numpy(values).unsqueeze(0).to(self.device)
        mask = torch.ones((1, keypoints.shape[1]), dtype=torch.float32, device=self.device)
        with torch.inference_mode():
            _, logits = self.model(keypoints, mask)
            probabilities = torch.softmax(logits, dim=-1)[0]
            index = int(torch.argmax(probabilities).item())
            confidence = float(probabilities[index].item())
        return KArSL502Prediction(
            self.labels[index], confidence, confidence >= self.threshold, hand_frames, observed
        )


@lru_cache(maxsize=1)
def karsl502_recognizer() -> SignBartKArSL502Recognizer:
    return SignBartKArSL502Recognizer(Settings.load().visual_threshold)
